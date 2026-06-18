import os
import shutil
import calendar
from datetime import datetime, date
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
from openpyxl import load_workbook

# ============================================================
# CONFIGURACION GENERAL
# ============================================================

BASE_DIR = r"C:\Automatizaciones\Nomina"
PLANTILLA = os.path.join(BASE_DIR, "plantilla", "NOMINA.xlsx")
SALIDAS_DIR = os.path.join(BASE_DIR, "salidas")

HOJA_PRIMERA = "Nómina 01-15"
HOJA_SEGUNDA = "Nómina 16-30"

FILA_INICIO_EMPLEADOS = 18
COLUMNA_NOMBRE = "E"

COLUMNAS_HORAS = {
    "HED": "U",      # Horas extras diurnas
    "HEN": "W",      # Horas extras nocturnas
    "HRDF": "Y",     # Horas recargo dominical/festivo
    "RN": "AA",      # Recargos nocturnos
    "RND": "AC",     # Recargo nocturno dominical
    "HEDF": "AE",    # Horas extras dominicales/festivas
    "HENDF": "AG",   # Horas extras nocturnas dominicales/festivas
}

COLUMNAS_COMISIONES = {
    "Comision contratacion y R+R": "AI",
    "Comision creditos": "AJ",
    "Comision SINDRI": "AK",
    "Comision venta de oro y traslado": "AL",
    "Comision ROSETT": "AM",
    "Bonificacion por reemplazo": "AN",
    "Bonificacion variable por desplazamiento jefes de zona": "AO",
    "Bonificacion variable N joyerias jefes de zona": "AP",
    "Bonificacion bunker": "AQ",
}

MESES = {
    1: "ENERO",
    2: "FEBRERO",
    3: "MARZO",
    4: "ABRIL",
    5: "MAYO",
    6: "JUNIO",
    7: "JULIO",
    8: "AGOSTO",
    9: "SEPTIEMBRE",
    10: "OCTUBRE",
    11: "NOVIEMBRE",
    12: "DICIEMBRE",
}


def centrar_ventana(win, ancho, alto):
    win.update_idletasks()
    sw = win.winfo_screenwidth()
    sh = win.winfo_screenheight()
    x = max((sw - ancho) // 2, 0)
    y = max((sh - alto) // 2, 0)
    win.geometry(f"{ancho}x{alto}+{x}+{y}")


def convertir_numero(valor):
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(",", ".")
    if not texto:
        return 0.0
    return float(texto)


def formatear_numero(valor):
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


class NominaApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Nomina Simple - GAMAN")
        self.root.resizable(True, True)
        centrar_ventana(self.root, 1050, 720)

        self.archivo_actual = None
        self.empleados = []
        self.comisiones = {}  # {tipo_comision: {fila_excel: valor}}
        self.calendario_widgets = {}
        self.nocturno_a_combo = None
        self.nocturno_b_combo = None

        self.crear_interfaz()

    def crear_interfaz(self):
        main = ttk.Frame(self.root, padding=10)
        main.pack(fill="both", expand=True)

        config = ttk.LabelFrame(main, text="Configuracion", padding=10)
        config.pack(fill="x", pady=5)

        ttk.Label(config, text="Anio:").grid(row=0, column=0, sticky="w", padx=5, pady=3)
        self.anio_var = tk.StringVar(value=str(datetime.now().year))
        ttk.Entry(config, textvariable=self.anio_var, width=10).grid(row=0, column=1, padx=5, pady=3)

        ttk.Label(config, text="Mes:").grid(row=0, column=2, sticky="w", padx=5, pady=3)
        self.mes_combo = ttk.Combobox(config, values=list(MESES.keys()), state="readonly", width=10)
        self.mes_combo.set(datetime.now().month)
        self.mes_combo.grid(row=0, column=3, padx=5, pady=3)

        ttk.Label(config, text="Quincena:").grid(row=0, column=4, sticky="w", padx=5, pady=3)
        self.quincena_combo = ttk.Combobox(config, values=["Primera", "Segunda"], state="readonly", width=12)
        self.quincena_combo.set("Primera")
        self.quincena_combo.grid(row=0, column=5, padx=5, pady=3)

        ttk.Label(config, text="Inicio novedades:").grid(row=1, column=0, sticky="w", padx=5, pady=3)
        self.inicio_var = tk.StringVar(value="01")
        ttk.Entry(config, textvariable=self.inicio_var, width=10).grid(row=1, column=1, padx=5, pady=3)

        ttk.Label(config, text="Fin novedades:").grid(row=1, column=2, sticky="w", padx=5, pady=3)
        self.fin_var = tk.StringVar(value="15")
        ttk.Entry(config, textvariable=self.fin_var, width=10).grid(row=1, column=3, padx=5, pady=3)

        botones = ttk.Frame(main)
        botones.pack(fill="x", pady=10)

        ttk.Button(botones, text="1. Crear archivo de nomina", command=self.crear_archivo_nomina).pack(side="left", padx=4)
        ttk.Button(botones, text="2. Cargar empleados", command=self.cargar_empleados).pack(side="left", padx=4)
        ttk.Button(botones, text="3. Modulo nocturnos", command=self.abrir_modulo_nocturnos).pack(side="left", padx=4)
        ttk.Button(botones, text="4. Modulo comisiones", command=self.abrir_modulo_comisiones).pack(side="left", padx=4)
        ttk.Button(botones, text="5. GUARDAR NOMINA", command=self.guardar_nomina_excel).pack(side="left", padx=12)

        info = ttk.Label(main, text="Doble clic sobre un empleado para editar horas manuales.", foreground="gray")
        info.pack(anchor="w", pady=(0, 5))

        tabla_frame = ttk.LabelFrame(main, text="Empleados y novedades", padding=8)
        tabla_frame.pack(fill="both", expand=True)

        columnas = ["Fila", "Nombre", "HED", "HEN", "HRDF", "RN", "RND", "HEDF", "HENDF"]
        self.tree = ttk.Treeview(tabla_frame, columns=columnas, show="headings", height=18)

        anchos = {"Fila": 60, "Nombre": 300, "HED": 80, "HEN": 80, "HRDF": 80, "RN": 80, "RND": 80, "HEDF": 80, "HENDF": 80}

        for col in columnas:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=anchos[col], anchor="center" if col != "Nombre" else "w")

        scroll_y = ttk.Scrollbar(tabla_frame, orient="vertical", command=self.tree.yview)
        scroll_x = ttk.Scrollbar(tabla_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        scroll_y.grid(row=0, column=1, sticky="ns")
        scroll_x.grid(row=1, column=0, sticky="ew")
        tabla_frame.rowconfigure(0, weight=1)
        tabla_frame.columnconfigure(0, weight=1)
        self.tree.bind("<Double-1>", self.editar_empleado)

    def obtener_hoja(self):
        return HOJA_PRIMERA if self.quincena_combo.get() == "Primera" else HOJA_SEGUNDA

    def crear_archivo_nomina(self):
        if not os.path.exists(PLANTILLA):
            messagebox.showerror("Error", f"No se encontro la plantilla:\n{PLANTILLA}")
            return
        try:
            anio = int(self.anio_var.get())
            mes = int(self.mes_combo.get())
        except ValueError:
            messagebox.showerror("Error", "Anio y mes deben ser numericos.")
            return
        if mes not in MESES:
            messagebox.showerror("Error", "Mes invalido.")
            return
        os.makedirs(SALIDAS_DIR, exist_ok=True)
        quincena_txt = "PRIMERA_QUINCENA" if self.quincena_combo.get() == "Primera" else "SEGUNDA_QUINCENA"
        nombre_archivo = f"NOMINA_{quincena_txt}_{MESES[mes]}_{anio}.xlsx"
        ruta_salida = os.path.join(SALIDAS_DIR, nombre_archivo)
        if os.path.exists(ruta_salida):
            respuesta = messagebox.askyesno("Archivo existente", f"Ya existe:\n{ruta_salida}\n\nDeseas reemplazarlo?")
            if not respuesta:
                return
        try:
            shutil.copy2(PLANTILLA, ruta_salida)
            self.archivo_actual = ruta_salida
            self.empleados.clear()
            self.comisiones.clear()
            self.actualizar_tabla()
            messagebox.showinfo("Exito", f"Archivo creado correctamente:\n{ruta_salida}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo crear el archivo:\n{e}")

    def cargar_empleados(self):
        if not self.archivo_actual or not os.path.exists(self.archivo_actual):
            messagebox.showerror("Error", "Primero debes crear el archivo de nomina.")
            return
        hoja_nombre = self.obtener_hoja()
        try:
            wb = load_workbook(self.archivo_actual, data_only=False)
            if hoja_nombre not in wb.sheetnames:
                wb.close()
                messagebox.showerror("Error", f"No existe la hoja:\n{hoja_nombre}")
                return
            ws = wb[hoja_nombre]
            self.empleados.clear()
            fila = FILA_INICIO_EMPLEADOS
            vacias = 0
            while vacias < 5:
                nombre = ws[f"{COLUMNA_NOMBRE}{fila}"].value
                if nombre and str(nombre).strip():
                    empleado = {
                        "fila": fila,
                        "nombre": str(nombre).strip(),
                        "HED": convertir_numero(ws[f"U{fila}"].value),
                        "HEN": convertir_numero(ws[f"W{fila}"].value),
                        "HRDF": convertir_numero(ws[f"Y{fila}"].value),
                        "RN": convertir_numero(ws[f"AA{fila}"].value),
                        "RND": convertir_numero(ws[f"AC{fila}"].value),
                        "HEDF": convertir_numero(ws[f"AE{fila}"].value),
                        "HENDF": convertir_numero(ws[f"AG{fila}"].value),
                    }
                    self.empleados.append(empleado)
                    vacias = 0
                else:
                    vacias += 1
                fila += 1
            wb.close()
            self.actualizar_tabla()
            messagebox.showinfo("Exito", f"Empleados cargados: {len(self.empleados)}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudieron cargar empleados:\n{e}")

    def actualizar_tabla(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for emp in self.empleados:
            self.tree.insert("", "end", values=(
                emp["fila"], emp["nombre"], formatear_numero(emp["HED"]), formatear_numero(emp["HEN"]),
                formatear_numero(emp["HRDF"]), formatear_numero(emp["RN"]), formatear_numero(emp["RND"]),
                formatear_numero(emp["HEDF"]), formatear_numero(emp["HENDF"]),
            ))

    def editar_empleado(self, event):
        seleccion = self.tree.selection()
        if not seleccion:
            return
        idx = self.tree.index(seleccion[0])
        emp = self.empleados[idx]
        win = tk.Toplevel(self.root)
        win.title(f"Editar horas - {emp['nombre']}")
        win.resizable(False, False)
        centrar_ventana(win, 430, 360)
        win.grab_set()
        frame = ttk.Frame(win, padding=15)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text=emp["nombre"], font=("Arial", 11, "bold")).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))
        entries = {}
        campos = ["HED", "HEN", "HRDF", "RN", "RND", "HEDF", "HENDF"]
        for i, campo in enumerate(campos, start=1):
            ttk.Label(frame, text=campo).grid(row=i, column=0, sticky="w", padx=5, pady=4)
            var = tk.StringVar(value=formatear_numero(emp[campo]))
            ttk.Entry(frame, textvariable=var, width=18).grid(row=i, column=1, padx=5, pady=4)
            entries[campo] = var
        def guardar():
            try:
                for campo in campos:
                    emp[campo] = convertir_numero(entries[campo].get())
                self.actualizar_tabla()
                win.destroy()
            except Exception:
                messagebox.showerror("Error", "Todos los campos deben ser numericos.")
        ttk.Button(frame, text="Guardar", command=guardar).grid(row=9, column=0, pady=15)
        ttk.Button(frame, text="Cancelar", command=win.destroy).grid(row=9, column=1, pady=15)

    def abrir_modulo_nocturnos(self):
        if not self.empleados:
            messagebox.showerror("Error", "Primero debes cargar empleados.")
            return
        win = tk.Toplevel(self.root)
        win.title("Modulo Nocturnos")
        win.resizable(True, True)
        centrar_ventana(win, 900, 650)
        win.grab_set()
        frame = ttk.Frame(win, padding=10)
        frame.pack(fill="both", expand=True)
        empleados_nombres = [e["nombre"] for e in self.empleados]
        top = ttk.LabelFrame(frame, text="Seleccion de nocturnos", padding=10)
        top.pack(fill="x", pady=5)
        ttk.Label(top, text="Nocturno A:").grid(row=0, column=0, sticky="w", padx=5, pady=3)
        self.nocturno_a_combo = ttk.Combobox(top, values=empleados_nombres, state="readonly", width=45)
        self.nocturno_a_combo.grid(row=0, column=1, padx=5, pady=3)
        ttk.Label(top, text="Nocturno B:").grid(row=1, column=0, sticky="w", padx=5, pady=3)
        self.nocturno_b_combo = ttk.Combobox(top, values=empleados_nombres, state="readonly", width=45)
        self.nocturno_b_combo.grid(row=1, column=1, padx=5, pady=3)
        cal_frame = ttk.LabelFrame(frame, text="Calendario mensual", padding=10)
        cal_frame.pack(fill="both", expand=True, pady=8)
        try:
            anio = int(self.anio_var.get())
            mes = int(self.mes_combo.get())
        except ValueError:
            messagebox.showerror("Error", "Anio o mes invalido.")
            win.destroy()
            return
        headers = ["Lun", "Mar", "Mie", "Jue", "Vie", "Sab", "Dom"]
        for c, h in enumerate(headers):
            ttk.Label(cal_frame, text=h, font=("Arial", 10, "bold")).grid(row=0, column=c, padx=4, pady=3)
        self.calendario_widgets = {}
        cal = calendar.monthcalendar(anio, mes)
        for r, week in enumerate(cal, start=1):
            for c, day in enumerate(week):
                if day == 0:
                    continue
                dframe = ttk.Frame(cal_frame, relief="solid", borderwidth=1, padding=3)
                dframe.grid(row=r, column=c, sticky="nsew", padx=3, pady=3)
                ttk.Label(dframe, text=str(day), font=("Arial", 10, "bold")).pack(anchor="w")
                turno = tk.StringVar(value="Libre")
                ttk.Combobox(dframe, textvariable=turno, values=["Libre", "A", "B"], width=6, state="readonly").pack(anchor="w")
                festivo = tk.BooleanVar(value=False)
                descanso = tk.BooleanVar(value=False)
                ttk.Checkbutton(dframe, text="Festivo", variable=festivo).pack(anchor="w")
                ttk.Checkbutton(dframe, text="Descanso", variable=descanso).pack(anchor="w")
                self.calendario_widgets[day] = {"turno": turno, "festivo": festivo, "descanso": descanso}
        for c in range(7):
            cal_frame.columnconfigure(c, weight=1)
        acciones = ttk.Frame(frame)
        acciones.pack(fill="x", pady=8)
        ttk.Button(acciones, text="Calcular y aplicar nocturnos", command=lambda: self.calcular_y_aplicar_nocturnos(win)).pack(side="left", padx=5)
        ttk.Button(acciones, text="Cerrar", command=win.destroy).pack(side="left", padx=5)

    def calcular_y_aplicar_nocturnos(self, win):
        nocturno_a = self.nocturno_a_combo.get().strip()
        nocturno_b = self.nocturno_b_combo.get().strip()
        if not nocturno_a or not nocturno_b:
            messagebox.showerror("Error", "Debes seleccionar Nocturno A y Nocturno B.")
            return
        if nocturno_a == nocturno_b:
            messagebox.showerror("Error", "Nocturno A y Nocturno B no pueden ser el mismo empleado.")
            return
        try:
            anio = int(self.anio_var.get())
            mes = int(self.mes_combo.get())
        except ValueError:
            messagebox.showerror("Error", "Anio o mes invalido.")
            return
        
        # Reiniciar solo las horas de los empleados seleccionados como nocturnos
        for emp in self.empleados:
            if emp["nombre"] in (nocturno_a, nocturno_b):
                for campo in COLUMNAS_HORAS.keys():
                    emp[campo] = 0.0
        
        acumulados = {nocturno_a: {campo: 0.0 for campo in COLUMNAS_HORAS.keys()}, 
                      nocturno_b: {campo: 0.0 for campo in COLUMNAS_HORAS.keys()}}
        
        # Recopilar dias trabajados por empleado
        dias_por_empleado = {nocturno_a: [], nocturno_b: []}
        festivo_descanso = {nocturno_a: [], nocturno_b: []}
        
        for day, widgets in self.calendario_widgets.items():
            turno = widgets["turno"].get()
            festivo = widgets["festivo"].get()
            descanso = widgets["descanso"].get()
            if turno == "Libre":
                continue
            empleado = nocturno_a if turno == "A" else nocturno_b
            if festivo and descanso:
                festivo_descanso[empleado].append(day)
                continue
            dias_por_empleado[empleado].append(day)
        
        # Procesar Festivo + Descanso (prioridad maxima, segun tabla PDF)
        for empleado, dias in festivo_descanso.items():
            for day in dias:
                # Valores exactos de tabla PDF
                acumulados[empleado]["HED"] += 2.5
                acumulados[empleado]["HEN"] += 4.5
                acumulados[empleado]["HRDF"] += 10.5
                acumulados[empleado]["HENDF"] += 5
        
        # Procesar dias normales por empleado
        for empleado, dias in dias_por_empleado.items():
            if not dias:
                continue
            
            dias_ordenados = sorted(set(dias))
            dias_procesados = set()
            
            # Detectar y procesar bloques de lunes a jueves (Nocturno 1)
            i = 0
            while i < len(dias_ordenados):
                day = dias_ordenados[i]
                if day in dias_procesados:
                    i += 1
                    continue
                
                fecha = date(anio, mes, day)
                dia_semana = fecha.weekday()  # 0=lunes, 6=domingo
                
                # Si es lunes (0), martes (1), miercoles (2) o jueves (3)
                if dia_semana <= 3:
                    # Iniciar bloque de lunes a jueves consecutivos
                    bloque = [day]
                    j = i + 1
                    while j < len(dias_ordenados):
                        next_day = dias_ordenados[j]
                        if next_day == day + 1:
                            next_fecha = date(anio, mes, next_day)
                            next_dia_semana = next_fecha.weekday()
                            if next_dia_semana <= 3:  # Sigue siendo lunes-jueves
                                bloque.append(next_day)
                                day = next_day
                                j += 1
                            else:
                                break
                        else:
                            break
                    
                    # Aplicar valores por cada dia del bloque (Nocturno 1)
                    for d in bloque:
                        acumulados[empleado]["HED"] += 2
                        acumulados[empleado]["RN"] += 9.5
                        dias_procesados.add(d)
                    
                    i = j
                else:
                    i += 1
            
            # Detectar y procesar bloques viernes-sabado-domingo (Nocturno 2)
            i = 0
            while i < len(dias_ordenados):
                day = dias_ordenados[i]
                if day in dias_procesados:
                    i += 1
                    continue
                
                fecha = date(anio, mes, day)
                dia_semana = fecha.weekday()
                
                # Si es viernes (4)
                if dia_semana == 4:
                    sab = day + 1
                    dom = day + 2
                    
                    # Verificar si sabado y domingo estan en la lista del MISMO empleado
                    if sab in dias_ordenados and dom in dias_ordenados:
                        try:
                            fecha_sab = date(anio, mes, sab)
                            fecha_dom = date(anio, mes, dom)
                            if fecha_sab.weekday() == 5 and fecha_dom.weekday() == 6:
                                # Bloque completo encontrado (Nocturno 2)
                                acumulados[empleado]["RN"] += 29
                                acumulados[empleado]["HEDF"] += 7.5
                                acumulados[empleado]["HENDF"] += 0.5
                                dias_procesados.update([day, sab, dom])
                                i += 3
                                continue
                        except ValueError:
                            pass
                    # Si no forma bloque completo, no hacer nada (no advertir, no calcular)
                
                # Si es sabado (5) o domingo (6) y no fue procesado
                # No hacer nada (no advertir, no calcular)
                
                i += 1
        
        # Aplicar valores a la tabla principal
        for emp in self.empleados:
            if emp["nombre"] in acumulados:
                for campo, valor in acumulados[emp["nombre"]].items():
                    emp[campo] = valor
        
        self.actualizar_tabla()
        
        # Mostrar resumen final limpio sin advertencias
        lineas = []
        for empleado, vals in acumulados.items():
            lineas.append(empleado)
            for campo in ["HED", "HEN", "HRDF", "RN", "RND", "HEDF", "HENDF"]:
                if vals[campo]:
                    lineas.append(f"  {campo}: {vals[campo]}")
            lineas.append("")
        
        messagebox.showinfo("Resumen nocturnos", "\n".join(lineas))
        win.destroy()

    def abrir_modulo_comisiones(self):
        if not self.empleados:
            messagebox.showerror("Error", "Primero debes cargar empleados.")
            return
        win = tk.Toplevel(self.root)
        win.title("Modulo Comisiones")
        win.resizable(True, True)
        centrar_ventana(win, 760, 600)
        win.grab_set()
        frame = ttk.Frame(win, padding=10)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="Tipo de comision:").grid(row=0, column=0, sticky="w", padx=5, pady=5)
        tipo_var = tk.StringVar(value=list(COLUMNAS_COMISIONES.keys())[0])
        tipo_combo = ttk.Combobox(frame, textvariable=tipo_var, values=list(COLUMNAS_COMISIONES.keys()), state="readonly", width=45)
        tipo_combo.grid(row=0, column=1, sticky="ew", padx=5, pady=5)
        ttk.Label(frame, text="Valor total:").grid(row=1, column=0, sticky="w", padx=5, pady=5)
        total_var = tk.StringVar(value="0")
        ttk.Entry(frame, textvariable=total_var, width=20).grid(row=1, column=1, sticky="w", padx=5, pady=5)
        repartir_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Repartir en partes iguales", variable=repartir_var).grid(row=2, column=0, columnspan=2, sticky="w", padx=5, pady=5)
        lista_frame = ttk.LabelFrame(frame, text="Empleados", padding=8)
        lista_frame.grid(row=3, column=0, columnspan=2, sticky="nsew", pady=8)
        frame.rowconfigure(3, weight=1)
        frame.columnconfigure(1, weight=1)
        lista_frame.rowconfigure(0, weight=1)
        lista_frame.columnconfigure(0, weight=1)
        cols = ["Seleccion", "Fila", "Nombre", "Valor manual"]
        tree = ttk.Treeview(lista_frame, columns=cols, show="headings", height=12)
        for col in cols:
            tree.heading(col, text=col)
        tree.column("Seleccion", width=80, anchor="center")
        tree.column("Fila", width=60, anchor="center")
        tree.column("Nombre", width=280, anchor="w")
        tree.column("Valor manual", width=120, anchor="center")
        scroll = ttk.Scrollbar(lista_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        scroll.grid(row=0, column=1, sticky="ns")
        seleccionados = {}
        for emp in self.empleados:
            item = tree.insert("", "end", values=(" ", emp["fila"], emp["nombre"], "0"))
            seleccionados[item] = False
        def toggle(event):
            item = tree.identify_row(event.y)
            col = tree.identify_column(event.x)
            if not item or col == "#4":
                return
            vals = list(tree.item(item, "values"))
            seleccionados[item] = not seleccionados[item]
            vals[0] = "X" if seleccionados[item] else " "
            tree.item(item, values=vals)
        def editar_valor_manual(event):
            item = tree.identify_row(event.y)
            col = tree.identify_column(event.x)
            if not item or col != "#4":
                return
            vals = list(tree.item(item, "values"))
            actual = vals[3]
            nuevo = simpledialog.askstring("Valor manual", f"Valor para {vals[2]}:", initialvalue=actual, parent=win)
            if nuevo is None:
                return
            try:
                valor = convertir_numero(nuevo)
            except Exception:
                messagebox.showerror("Error", "Valor invalido.")
                return
            vals[3] = formatear_numero(valor)
            tree.item(item, values=vals)
        tree.bind("<Button-1>", toggle)
        tree.bind("<Double-1>", editar_valor_manual)
        def guardar_comision():
            tipo = tipo_var.get()
            try:
                total = convertir_numero(total_var.get())
            except Exception:
                messagebox.showerror("Error", "Valor total invalido.")
                return
            if total <= 0:
                messagebox.showerror("Error", "El valor total debe ser mayor a cero.")
                return
            distribucion = {}
            if repartir_var.get():
                items_sel = [item for item, sel in seleccionados.items() if sel]
                if not items_sel:
                    messagebox.showerror("Error", "Selecciona al menos un empleado.")
                    return
                valor_individual = total / len(items_sel)
                for item in items_sel:
                    vals = tree.item(item, "values")
                    fila = int(vals[1])
                    distribucion[fila] = valor_individual
            else:
                suma = 0.0
                for item in tree.get_children():
                    vals = tree.item(item, "values")
                    fila = int(vals[1])
                    valor = convertir_numero(vals[3])
                    if valor > 0:
                        distribucion[fila] = valor
                        suma += valor
                if not distribucion:
                    messagebox.showerror("Error", "Ingresa al menos un valor manual.")
                    return
                if abs(suma - total) > 0.01:
                    diferencia = total - suma
                    if diferencia > 0:
                        messagebox.showerror("Error", f"Falta asignar ${diferencia:,.2f}")
                    else:
                        messagebox.showerror("Error", f"Sobra ${abs(diferencia):,.2f}")
                    return
            self.comisiones[tipo] = distribucion
            # Mostrar resumen detallado
            total_asignado = sum(distribucion.values())
            empleados_asignados = []
            for fila, valor in distribucion.items():
                for emp in self.empleados:
                    if emp["fila"] == fila:
                        empleados_asignados.append(f"{emp['nombre']}: ${valor:,.2f}")
                        break
            resumen = f"Tipo de comision: {tipo}\n"
            resumen += f"Total asignado: ${total_asignado:,.2f}\n"
            resumen += f"Empleados asignados ({len(distribucion)}):\n"
            resumen += "\n".join(empleados_asignados)
            messagebox.showinfo("Comision guardada", resumen)
            win.destroy()
        botones = ttk.Frame(frame)
        botones.grid(row=4, column=0, columnspan=2, pady=8)
        ttk.Button(botones, text="Guardar comision en memoria", command=guardar_comision).pack(side="left", padx=5)
        ttk.Button(botones, text="Cerrar", command=win.destroy).pack(side="left", padx=5)

    def validar_nomina(self):
        errores_criticos = []
        advertencias = []
        
        # Validar archivo_actual
        if not self.archivo_actual:
            errores_criticos.append("El archivo actual esta vacio o no existe.")
        
        # Validar empleados cargados
        if not self.empleados:
            errores_criticos.append("No hay empleados cargados.")
        
        # Validar horas negativas
        for emp in self.empleados:
            for campo in COLUMNAS_HORAS.keys():
                if emp[campo] < 0:
                    errores_criticos.append(f"Empleado {emp['nombre']} tiene horas negativas en {campo}: {emp[campo]}")
        
        # Calcular totales para advertencias
        empleados_con_horas = []
        total_horas = 0
        total_comisiones = 0
        
        for emp in self.empleados:
            suma_horas = sum(emp[campo] for campo in COLUMNAS_HORAS.keys())
            if suma_horas > 0:
                empleados_con_horas.append(emp)
                total_horas += suma_horas
            
            if suma_horas > 60:
                advertencias.append(f"Empleado {emp['nombre']} tiene total de horas mayor a 60: {suma_horas}")
            elif suma_horas > 40:
                advertencias.append(f"Empleado {emp['nombre']} tiene total de horas mayor a 40: {suma_horas}")
        
        # Validar comisiones individuales
        for tipo, distribucion in self.comisiones.items():
            for fila, valor in distribucion.items():
                total_comisiones += valor
                if valor > 5000000:
                    for emp in self.empleados:
                        if emp["fila"] == fila:
                            advertencias.append(f"Empleado {emp['nombre']} tiene comision individual mayor a 5,000,000: ${valor:,.2f}")
                            break
        
        # Validar si hay comisiones cargadas
        if not self.comisiones:
            advertencias.append("No hay comisiones cargadas.")
        
        # Validar si hay empleados con novedades
        if not empleados_con_horas:
            advertencias.append("No hay empleados con novedades (horas).")
        
        # Calcular resumen
        resumen = {
            "empleados_con_horas": empleados_con_horas,
            "total_horas": total_horas,
            "total_comisiones": total_comisiones,
            "tipos_comision": len(self.comisiones)
        }
        
        return errores_criticos, advertencias, resumen
    
    def guardar_nomina_excel(self):
        if not self.archivo_actual or not os.path.exists(self.archivo_actual):
            messagebox.showerror("Error", "Primero debes crear el archivo de nomina.")
            return
        if not self.empleados:
            messagebox.showerror("Error", "Primero debes cargar empleados.")
            return
        
        # Ejecutar validacion antes de guardar
        errores_criticos, advertencias, resumen = self.validar_nomina()
        
        # Si hay errores criticos, mostrar y no guardar
        if errores_criticos:
            mensaje = "ERRORES CRITICOS encontrados:\n\n"
            mensaje += "\n".join(f"- {error}" for error in errores_criticos)
            messagebox.showerror("Error de validacion", mensaje)
            return
        
        # Si hay advertencias, mostrar resumen y preguntar
        if advertencias:
            mensaje_resumen = "RESUMEN DE NOMINA:\n\n"
            mensaje_resumen += f"Empleados con horas: {len(resumen['empleados_con_horas'])}\n"
            mensaje_resumen += f"Total horas: {resumen['total_horas']}\n"
            mensaje_resumen += f"Total comisiones: ${resumen['total_comisiones']:,.2f}\n"
            mensaje_resumen += f"Tipos de comision: {resumen['tipos_comision']}\n\n"
            
            mensaje_advertencias = "ADVERTENCIAS:\n\n"
            mensaje_advertencias += "\n".join(f"- {adv}" for adv in advertencias)
            
            mensaje_completo = mensaje_resumen + mensaje_advertencias + "\n\nDeseas continuar con el guardado?"
            
            respuesta = messagebox.askyesno("Advertencias de validacion", mensaje_completo)
            if not respuesta:
                return
        
        # Si no hay errores ni advertencias, o si el usuario confirmo, guardar
        hoja_nombre = self.obtener_hoja()
        try:
            wb = load_workbook(self.archivo_actual, data_only=False)
            if hoja_nombre not in wb.sheetnames:
                wb.close()
                messagebox.showerror("Error", f"No existe la hoja:\n{hoja_nombre}")
                return
            ws = wb[hoja_nombre]
            for emp in self.empleados:
                fila = emp["fila"]
                for campo, columna in COLUMNAS_HORAS.items():
                    ws[f"{columna}{fila}"] = emp[campo]
            for tipo, distribucion in self.comisiones.items():
                columna = COLUMNAS_COMISIONES.get(tipo)
                if not columna:
                    continue
                for emp in self.empleados:
                    ws[f"{columna}{emp['fila']}"] = 0
                for fila, valor in distribucion.items():
                    ws[f"{columna}{fila}"] = valor
            wb.save(self.archivo_actual)
            wb.close()
            messagebox.showinfo("Exito", f"Nomina guardada correctamente:\n{self.archivo_actual}")
        except PermissionError:
            messagebox.showerror("Error", "No se pudo guardar el archivo.\n\nCierra el Excel si lo tienes abierto y vuelve a intentar.")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo guardar la nomina:\n{e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = NominaApp(root)
    root.mainloop()
