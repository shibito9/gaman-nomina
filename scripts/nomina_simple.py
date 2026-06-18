import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
from tkinter import Toplevel
import openpyxl
from openpyxl import load_workbook
import shutil
from datetime import datetime, date
import calendar
import os

# Rutas
BASE_DIR = r"C:\Automatizaciones\Nomina"
PLANTILLA = os.path.join(BASE_DIR, "plantilla", "NOMINA.xlsx")
SALIDAS_DIR = os.path.join(BASE_DIR, "salidas")

# Variable global principal
archivo_actual = None

# Datos en memoria
empleados = []  # Lista de diccionarios con datos de empleados
comisiones = {}  # Diccionario de comisiones por empleado

# Mapeo de columnas de horas
COLUMNAS_HORAS = {
    "HED": "U",
    "HEN": "W",
    "HRDF": "Y",
    "RN": "AA",
    "RND": "AC",
    "HEDF": "AE",
    "HENDF": "AG"
}

# Mapeo de columnas de comisiones
COLUMNAS_COMISIONES = {
    "Comisión contratación y R+R": "AI",
    "Comisión créditos": "AJ",
    "Comisión SINDRI": "AK",
    "Comisión venta de oro y traslado": "AL",
    "Comisión ROSETT": "AM",
    "Bonificación por reemplazo": "AN",
    "Bonificación variable por desplazamiento jefes de zona": "AO",
    "Bonificación variable N° joyerías jefes de zona": "AP",
    "Bonificación bunker": "AQ"
}

# Nombres de meses
MESES_NOMBRE = {
    1: "ENERO",
    2: "FEBRERO",
    3: "MARZO",
    4: "ABRIL",
    5: "MAYO",
    6: "JUNIO",
    7: "JULIO",
    8: "AGOSTO",
    9: "SEPTIEMBRE",
    10: "OCTUBRE",
    11: "NOVIEMBRE",
    12: "DICIEMBRE"
}


class NominaSimple:
    def __init__(self, root):
        self.root = root
        self.root.title("Nómina Simple")
        self.root.geometry("900x700")
        
        self.nocturno_a = None
        self.nocturno_b = None
        self.calendario_data = {}  # {dia: {"turno": "A"/"B"/"Libre", "festivo": bool, "descanso": bool}}
        
        self.crear_interfaz()
    
    def crear_interfaz(self):
        # Frame principal
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Frame de configuración
        config_frame = ttk.LabelFrame(main_frame, text="Configuración", padding="10")
        config_frame.pack(fill=tk.X, pady=5)
        
        # Año
        ttk.Label(config_frame, text="Año:").grid(row=0, column=0, sticky=tk.W, padx=5)
        self.anio_var = tk.StringVar(value=str(datetime.now().year))
        ttk.Entry(config_frame, textvariable=self.anio_var, width=10).grid(row=0, column=1, padx=5)
        
        # Mes
        ttk.Label(config_frame, text="Mes:").grid(row=0, column=2, sticky=tk.W, padx=5)
        self.mes_combo = ttk.Combobox(config_frame, values=list(range(1, 13)), width=10)
        self.mes_combo.set(datetime.now().month)
        self.mes_combo.grid(row=0, column=3, padx=5)
        
        # Quincena
        ttk.Label(config_frame, text="Quincena:").grid(row=0, column=4, sticky=tk.W, padx=5)
        self.quincena_combo = ttk.Combobox(config_frame, values=["Primera", "Segunda"], width=15)
        self.quincena_combo.set("Primera")
        self.quincena_combo.grid(row=0, column=5, padx=5)
        
        # Inicio novedades
        ttk.Label(config_frame, text="Inicio Novedades:").grid(row=1, column=0, sticky=tk.W, padx=5)
        self.inicio_novedades_var = tk.StringVar(value="01")
        ttk.Entry(config_frame, textvariable=self.inicio_novedades_var, width=10).grid(row=1, column=1, padx=5)
        
        # Fin novedades
        ttk.Label(config_frame, text="Fin Novedades:").grid(row=1, column=2, sticky=tk.W, padx=5)
        self.fin_novedades_var = tk.StringVar(value="15")
        ttk.Entry(config_frame, textvariable=self.fin_novedades_var, width=10).grid(row=1, column=3, padx=5)
        
        # Frame de botones principales
        botones_frame = ttk.Frame(main_frame)
        botones_frame.pack(fill=tk.X, pady=10)
        
        ttk.Button(botones_frame, text="Crear archivo de nómina", command=self.crear_archivo_nomina).pack(side=tk.LEFT, padx=5)
        ttk.Button(botones_frame, text="Cargar empleados", command=self.cargar_empleados).pack(side=tk.LEFT, padx=5)
        ttk.Button(botones_frame, text="Módulo Nocturnos", command=self.abrir_modulo_nocturnos).pack(side=tk.LEFT, padx=5)
        ttk.Button(botones_frame, text="Módulo Comisiones", command=self.abrir_modulo_comisiones).pack(side=tk.LEFT, padx=5)
        ttk.Button(botones_frame, text="Guardar todo en Excel", command=self.guardar_todo_excel).pack(side=tk.LEFT, padx=5)
        
        # Frame de tabla de empleados
        tabla_frame = ttk.LabelFrame(main_frame, text="Empleados", padding="10")
        tabla_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # Treeview para empleados
        columns = ["Nombre", "HED", "HEN", "HRDF", "RN", "RND", "HEDF", "HENDF"]
        self.tree = ttk.Treeview(tabla_frame, columns=columns, show="headings")
        
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=100)
        
        scrollbar = ttk.Scrollbar(tabla_frame, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Doble clic para editar
        self.tree.bind("<Double-1>", self.editar_empleado)
    
    def crear_archivo_nomina(self):
        global archivo_actual
        
        try:
            # Verificar que la plantilla existe
            if not os.path.exists(PLANTILLA):
                messagebox.showerror("Error", f"No se encuentra la plantilla: {PLANTILLA}")
                return
            
            # Crear directorio de salidas si no existe
            os.makedirs(SALIDAS_DIR, exist_ok=True)
            
            # Obtener valores
            anio = self.anio_var.get()
            mes = self.mes_combo.get()
            quincena = self.quincena_combo.get()
            
            # Convertir mes a entero
            try:
                mes_numero = int(mes)
            except ValueError:
                messagebox.showerror("Error", "El mes debe ser un número.")
                return
            
            # Generar nombre de archivo
            quincena_num = "01-15" if quincena == "Primera" else "16-30"
            nombre_mes = MESES_NOMBRE[mes_numero]
            nombre_archivo = f"NOMINA_{anio}_{nombre_mes}_{quincena_num}.xlsx"
            ruta_destino = os.path.join(SALIDAS_DIR, nombre_archivo)
            
            # Copiar plantilla
            shutil.copy2(PLANTILLA, ruta_destino)
            
            # Guardar en variable global
            archivo_actual = ruta_destino
            
            messagebox.showinfo("Éxito", f"Archivo creado:\n{ruta_destino}")
            
        except Exception as e:
            messagebox.showerror("Error", f"Error al crear archivo: {str(e)}")
    
    def cargar_empleados(self):
        global archivo_actual, empleados
        
        if not archivo_actual or not os.path.exists(archivo_actual):
            messagebox.showerror("Error", "Primero debe crear un archivo de nómina.")
            return
        
        try:
            # Determinar hoja según quincena
            quincena = self.quincena_combo.get()
            nombre_hoja = "Nómina 01-15" if quincena == "Primera" else "Nómina 16-30"
            
            # Cargar workbook
            wb = load_workbook(archivo_actual)
            
            if nombre_hoja not in wb.sheetnames:
                messagebox.showerror("Error", f"No existe la hoja '{nombre_hoja}' en el archivo.")
                wb.close()
                return
            
            ws = wb[nombre_hoja]
            
            # Leer empleados desde E18 hacia abajo hasta 5 filas vacías
            empleados = []
            fila = 18
            filas_vacias = 0
            
            while filas_vacias < 5:
                nombre = ws[f"E{fila}"].value
                if nombre and str(nombre).strip():
                    # Leer horas actuales
                    empleado = {
                        "fila": fila,
                        "nombre": str(nombre).strip(),
                        "HED": ws[f"U{fila}"].value or 0,
                        "HEN": ws[f"W{fila}"].value or 0,
                        "HRDF": ws[f"Y{fila}"].value or 0,
                        "RN": ws[f"AA{fila}"].value or 0,
                        "RND": ws[f"AC{fila}"].value or 0,
                        "HEDF": ws[f"AE{fila}"].value or 0,
                        "HENDF": ws[f"AG{fila}"].value or 0
                    }
                    empleados.append(empleado)
                    filas_vacias = 0
                else:
                    filas_vacias += 1
                fila += 1
            
            wb.close()
            
            # Actualizar tabla
            self.actualizar_tabla()
            
            messagebox.showinfo("Éxito", f"Se cargaron {len(empleados)} empleados.")
            
        except Exception as e:
            messagebox.showerror("Error", f"Error al cargar empleados: {str(e)}")
    
    def actualizar_tabla(self):
        # Limpiar tabla
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        # Llenar tabla
        for emp in empleados:
            self.tree.insert("", tk.END, values=(
                emp["nombre"],
                emp["HED"],
                emp["HEN"],
                emp["HRDF"],
                emp["RN"],
                emp["RND"],
                emp["HEDF"],
                emp["HENDF"]
            ))
    
    def editar_empleado(self, event):
        selection = self.tree.selection()
        if not selection:
            return
        
        item = selection[0]
        idx = self.tree.index(item)
        empleado = empleados[idx]
        
        # Crear ventana de edición
        edit_window = Toplevel(self.root)
        edit_window.title(f"Editar: {empleado['nombre']}")
        edit_window.geometry("400x300")
        
        frame = ttk.Frame(edit_window, padding="10")
        frame.pack(fill=tk.BOTH, expand=True)
        
        entries = {}
        row = 0
        for campo in ["HED", "HEN", "HRDF", "RN", "RND", "HEDF", "HENDF"]:
            ttk.Label(frame, text=f"{campo}:").grid(row=row, column=0, sticky=tk.W, padx=5, pady=5)
            var = tk.StringVar(value=str(empleado[campo]))
            entry = ttk.Entry(frame, textvariable=var, width=15)
            entry.grid(row=row, column=1, padx=5, pady=5)
            entries[campo] = var
            row += 1
        
        def guardar():
            try:
                for campo, var in entries.items():
                    empleado[campo] = float(var.get())
                self.actualizar_tabla()
                edit_window.destroy()
            except ValueError:
                messagebox.showerror("Error", "Ingrese valores numéricos válidos.")
        
        ttk.Button(frame, text="Guardar", command=guardar).grid(row=row, column=0, columnspan=2, pady=10)
    
    def abrir_modulo_nocturnos(self):
        if not empleados:
            messagebox.showerror("Error", "Primero debe cargar empleados.")
            return
        
        nocturno_window = Toplevel(self.root)
        nocturno_window.title("Módulo Nocturnos")
        nocturno_window.geometry("800x600")
        
        frame = ttk.Frame(nocturno_window, padding="10")
        frame.pack(fill=tk.BOTH, expand=True)
        
        # Selección de nocturnos
        ttk.Label(frame, text="Nocturno A:").grid(row=0, column=0, sticky=tk.W, padx=5)
        nombres_empleados = [emp["nombre"] for emp in empleados]
        self.nocturno_a_combo = ttk.Combobox(frame, values=nombres_empleados, width=30)
        self.nocturno_a_combo.grid(row=0, column=1, padx=5)
        
        ttk.Label(frame, text="Nocturno B:").grid(row=0, column=2, sticky=tk.W, padx=5)
        self.nocturno_b_combo = ttk.Combobox(frame, values=nombres_empleados, width=30)
        self.nocturno_b_combo.grid(row=0, column=3, padx=5)
        
        # Calendario
        ttk.Label(frame, text="Calendario:").grid(row=1, column=0, columnspan=4, sticky=tk.W, pady=10)
        
        calendario_frame = ttk.Frame(frame)
        calendario_frame.grid(row=2, column=0, columnspan=4, sticky=tk.W)
        
        # Obtener mes y año
        try:
            anio = int(self.anio_var.get())
            mes = int(self.mes_combo.get())
        except:
            anio = datetime.now().year
            mes = datetime.now().month
        
        # Crear calendario
        cal = calendar.monthcalendar(anio, mes)
        
        # Headers
        headers = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
        for col, header in enumerate(headers):
            ttk.Label(calendario_frame, text=header, width=8).grid(row=0, column=col)
        
        # Días
        self.calendario_widgets = {}
        for row_week, week in enumerate(cal):
            for col_day, day in enumerate(week):
                if day == 0:
                    continue
                
                day_frame = ttk.Frame(calendario_frame, borderwidth=1, relief="solid")
                day_frame.grid(row=row_week + 1, column=col_day, padx=2, pady=2)
                
                ttk.Label(day_frame, text=str(day)).pack()
                
                # Combo para turno
                turno_var = tk.StringVar(value="Libre")
                turno_combo = ttk.Combobox(day_frame, textvariable=turno_var, values=["Libre", "A", "B"], width=5, state="readonly")
                turno_combo.pack()
                
                # Check festivo
                festivo_var = tk.BooleanVar(value=False)
                ttk.Checkbutton(day_frame, text="Festivo", variable=festivo_var).pack()
                
                # Check descanso trabajado
                descanso_var = tk.BooleanVar(value=False)
                ttk.Checkbutton(day_frame, text="Descanso", variable=descanso_var).pack()
                
                self.calendario_widgets[day] = {
                    "turno": turno_var,
                    "festivo": festivo_var,
                    "descanso": descanso_var
                }
        
        # Botón aplicar
        ttk.Button(frame, text="Aplicar nocturnos a tabla", command=lambda: self.aplicar_nocturnos(nocturno_window)).grid(row=3, column=0, columnspan=4, pady=10)
    
    def aplicar_nocturnos(self, window):
        nocturno_a = self.nocturno_a_combo.get()
        nocturno_b = self.nocturno_b_combo.get()
        
        if not nocturno_a or not nocturno_b:
            messagebox.showerror("Error", "Seleccione ambos nocturnos.")
            return
        
        # Obtener mes y año
        try:
            anio = int(self.anio_var.get())
            mes = int(self.mes_combo.get())
        except:
            messagebox.showerror("Error", "Año o mes inválido.")
            return
        
        print("\n" + "="*60)
        print("CÁLCULO DE NOCTURNOS")
        print("="*60)
        
        # Reiniciar acumulados (no acumular sobre cálculos anteriores)
        acumulados = {emp["nombre"]: {
            "HED": 0, "HEN": 0, "HRDF": 0, "RN": 0,
            "RND": 0, "HEDF": 0, "HENDF": 0
        } for emp in empleados}
        
        # Primero, recopilar todos los días trabajados por empleado
        dias_por_empleado = {nocturno_a: [], nocturno_b: []}
        festivos_descanso = {}  # {empleado: [dias]}
        
        for day, widgets in self.calendario_widgets.items():
            turno = widgets["turno"].get()
            festivo = widgets["festivo"].get()
            descanso = widgets["descanso"].get()
            
            if turno == "Libre":
                continue
            
            empleado = nocturno_a if turno == "A" else nocturno_b
            
            if festivo and descanso:
                if empleado not in festivos_descanso:
                    festivos_descanso[empleado] = []
                festivos_descanso[empleado].append(day)
            else:
                dias_por_empleado[empleado].append(day)
        
        # Procesar Festivo + Descanso (prioridad máxima)
        for empleado, dias in festivos_descanso.items():
            print(f"\n{empleado} - Festivos + Descanso:")
            for day in dias:
                print(f"  Día {day}: HED+=2.5, HEN+=4.5, HRDF+=10.5, HENDF+=5")
                acumulados[empleado]["HED"] += 2.5
                acumulados[empleado]["HEN"] += 4.5
                acumulados[empleado]["HRDF"] += 10.5
                acumulados[empleado]["HENDF"] += 5
        
        # Procesar bloques de fin de semana (viernes-sábado-domingo consecutivos)
        for empleado, dias in dias_por_empleado.items():
            if not dias:
                continue
            
            dias_ordenados = sorted(dias)
            print(f"\n{empleado} - Días trabajados (excluyendo festivo+descanso): {dias_ordenados}")
            
            # Detectar bloques viernes-sábado-domingo consecutivos
            i = 0
            dias_procesados = set()
            
            while i < len(dias_ordenados):
                day = dias_ordenados[i]
                fecha = date(anio, mes, day)
                dia_semana = fecha.weekday()  # 0=lunes, 6=domingo
                
                # Si es viernes (4), verificar si hay sábado (5) y domingo (6) consecutivos
                if dia_semana == 4:  # Viernes
                    # Verificar si el siguiente día es sábado y está en la lista
                    if i + 1 < len(dias_ordenados) and dias_ordenados[i + 1] == day + 1:
                        fecha_sab = date(anio, mes, day + 1)
                        if fecha_sab.weekday() == 5:  # Es sábado
                            # Verificar si el siguiente es domingo
                            if i + 2 < len(dias_ordenados) and dias_ordenados[i + 2] == day + 2:
                                fecha_dom = date(anio, mes, day + 2)
                                if fecha_dom.weekday() == 6:  # Es domingo
                                    # Bloque completo encontrado
                                    print(f"  Bloque viernes-sábado-domingo: días {day}, {day+1}, {day+2}")
                                    print(f"    RN+=29, HEDF+=7.5, HENDF+=0.5")
                                    acumulados[empleado]["RN"] += 29
                                    acumulados[empleado]["HEDF"] += 7.5
                                    acumulados[empleado]["HENDF"] += 0.5
                                    dias_procesados.add(day)
                                    dias_procesados.add(day + 1)
                                    dias_procesados.add(day + 2)
                                    i += 3
                                    continue
                
                i += 1
            
            # Procesar días restantes (no en bloques, no festivo+descanso)
            for day in dias_ordenados:
                if day in dias_procesados:
                    continue
                
                fecha = date(anio, mes, day)
                dia_semana = fecha.weekday()
                
                if dia_semana < 4:  # Lunes a jueves
                    print(f"  Día {day} (lunes-jueves): HED+=2, RN+=9.5")
                    acumulados[empleado]["HED"] += 2
                    acumulados[empleado]["RN"] += 9.5
                elif dia_semana >= 4:  # Viernes a domingo (no formó bloque)
                    print(f"  Día {day} (viernes-domingo, sin bloque): HED+=2, RN+=9.5")
                    acumulados[empleado]["HED"] += 2
                    acumulados[empleado]["RN"] += 9.5
        
        # Mostrar totales finales
        print("\n" + "="*60)
        print("TOTALES FINALES")
        print("="*60)
        for empleado in acumulados:
            print(f"\n{empleado}:")
            for campo, valor in acumulados[empleado].items():
                if valor > 0:
                    print(f"  {campo}: {valor}")
        
        # Reemplazar valores (no acumular sobre existentes)
        for emp in empleados:
            if emp["nombre"] in acumulados:
                for campo in acumulados[emp["nombre"]]:
                    emp[campo] = acumulados[emp["nombre"]][campo]
        
        self.actualizar_tabla()
        messagebox.showinfo("Éxito", "Nocturnos aplicados a la tabla.")
        window.destroy()
    
    def abrir_modulo_comisiones(self):
        if not empleados:
            messagebox.showerror("Error", "Primero debe cargar empleados.")
            return
        
        comision_window = Toplevel(self.root)
        comision_window.title("Módulo Comisiones")
        comision_window.geometry("600x400")
        
        frame = ttk.Frame(comision_window, padding="10")
        frame.pack(fill=tk.BOTH, expand=True)
        
        # Tipo de comisión
        ttk.Label(frame, text="Tipo de comisión:").grid(row=0, column=0, sticky=tk.W, padx=5)
        self.tipo_comision_combo = ttk.Combobox(frame, values=list(COLUMNAS_COMISIONES.keys()), width=40)
        self.tipo_comision_combo.grid(row=0, column=1, padx=5)
        
        # Valor total
        ttk.Label(frame, text="Valor total:").grid(row=1, column=0, sticky=tk.W, padx=5)
        self.valor_total_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.valor_total_var, width=15).grid(row=1, column=1, sticky=tk.W, padx=5)
        
        # Repartir igual
        ttk.Label(frame, text="Repartir igual:").grid(row=2, column=0, sticky=tk.W, padx=5)
        self.repartir_igual_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, variable=self.repartir_igual_var).grid(row=2, column=1, sticky=tk.W, padx=5)
        
        # Frame para selección de empleados
        self.empleados_frame = ttk.LabelFrame(frame, text="Seleccionar empleados", padding="10")
        self.empleados_frame.grid(row=3, column=0, columnspan=2, sticky=tk.W, pady=10)
        
        self.empleados_vars = {}
        for i, emp in enumerate(empleados):
            var = tk.BooleanVar(value=False)
            ttk.Checkbutton(self.empleados_frame, text=emp["nombre"], variable=var).grid(row=i//3, column=i%3, sticky=tk.W, padx=5)
            self.empleados_vars[emp["nombre"]] = var
        
        # Frame para valores individuales
        self.valores_frame = ttk.LabelFrame(frame, text="Valores por empleado", padding="10")
        self.valores_frame.grid(row=4, column=0, columnspan=2, sticky=tk.W, pady=10)
        
        self.valores_vars = {}
        for i, emp in enumerate(empleados):
            ttk.Label(self.valores_frame, text=emp["nombre"]).grid(row=i, column=0, sticky=tk.W, padx=5)
            var = tk.StringVar(value="0")
            ttk.Entry(self.valores_frame, textvariable=var, width=10).grid(row=i, column=1, padx=5)
            self.valores_vars[emp["nombre"]] = var
        
        # Botón guardar
        ttk.Button(frame, text="Guardar comisión", command=lambda: self.guardar_comision(comision_window)).grid(row=5, column=0, columnspan=2, pady=10)
    
    def guardar_comision(self, window):
        global comisiones
        
        tipo = self.tipo_comision_combo.get()
        if not tipo:
            messagebox.showerror("Error", "Seleccione tipo de comisión.")
            return
        
        try:
            valor_total = float(self.valor_total_var.get())
        except ValueError:
            messagebox.showerror("Error", "Ingrese un valor total válido.")
            return
        
        repartir_igual = self.repartir_igual_var.get()
        
        if repartir_igual:
            # Seleccionar empleados
            seleccionados = [nombre for nombre, var in self.empleados_vars.items() if var.get()]
            if not seleccionados:
                messagebox.showerror("Error", "Seleccione al menos un empleado.")
                return
            
            # Dividir valor
            valor_por_empleado = valor_total / len(seleccionados)
            
            # Guardar
            for nombre in seleccionados:
                if nombre not in comisiones:
                    comisiones[nombre] = {}
                comisiones[nombre][tipo] = valor_por_empleado
        else:
            # Usar valores individuales
            for nombre, var in self.valores_vars.items():
                try:
                    valor = float(var.get())
                    if valor > 0:
                        if nombre not in comisiones:
                            comisiones[nombre] = {}
                        comisiones[nombre][tipo] = valor
                except ValueError:
                    continue
        
        messagebox.showinfo("Éxito", "Comisión guardada en memoria.")
        window.destroy()
    
    def guardar_todo_excel(self):
        global archivo_actual, empleados, comisiones
        
        if not archivo_actual or not os.path.exists(archivo_actual):
            messagebox.showerror("Error", "Primero debe crear un archivo de nómina.")
            return
        
        if not empleados:
            messagebox.showerror("Error", "No hay empleados cargados.")
            return
        
        try:
            # Determinar hoja según quincena
            quincena = self.quincena_combo.get()
            nombre_hoja = "Nómina 01-15" if quincena == "Primera" else "Nómina 16-30"
            
            # Cargar workbook
            wb = load_workbook(archivo_actual)
            
            if nombre_hoja not in wb.sheetnames:
                messagebox.showerror("Error", f"No existe la hoja '{nombre_hoja}' en el archivo.")
                wb.close()
                return
            
            ws = wb[nombre_hoja]
            
            # Escribir horas
            for emp in empleados:
                fila = emp["fila"]
                ws[f"U{fila}"] = emp["HED"]
                ws[f"W{fila}"] = emp["HEN"]
                ws[f"Y{fila}"] = emp["HRDF"]
                ws[f"AA{fila}"] = emp["RN"]
                ws[f"AC{fila}"] = emp["RND"]
                ws[f"AE{fila}"] = emp["HEDF"]
                ws[f"AG{fila}"] = emp["HENDF"]
            
            # Escribir comisiones
            for emp in empleados:
                fila = emp["fila"]
                nombre = emp["nombre"]
                
                if nombre in comisiones:
                    for tipo_comision, valor in comisiones[nombre].items():
                        if tipo_comision in COLUMNAS_COMISIONES:
                            columna = COLUMNAS_COMISIONES[tipo_comision]
                            ws[f"{columna}{fila}"] = valor
            
            # Guardar y cerrar
            wb.save(archivo_actual)
            wb.close()
            
            messagebox.showinfo("Éxito", f"Datos guardados en:\n{archivo_actual}")
            
        except Exception as e:
            messagebox.showerror("Error", f"Error al guardar en Excel: {str(e)}")


if __name__ == "__main__":
    root = tk.Tk()
    app = NominaSimple(root)
    root.mainloop()
